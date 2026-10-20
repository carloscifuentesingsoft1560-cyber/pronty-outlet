import csv
import io
import re
import unicodedata

from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, cast
from uuid import uuid4

from django import forms
from django.db import transaction
from django.utils import timezone
from django.utils.text import slugify

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from .models import (
    Brand,
    Category,
    InventoryMovement,
    Product,
)


# ============================================================
# CONFIGURACIÓN GENERAL
# ============================================================

MAX_FILE_SIZE = 15 * 1024 * 1024
MAX_ROWS = 20000


ALLOWED_EXTENSIONS = {
    '.xlsx',
    '.csv',
}


CANONICAL_COLUMNS = [
    'sku',
    'nombre',
    'categoria',
    'marca',
    'descripcion_corta',
    'descripcion',
    'precio_detal',
    'precio_mayorista',
    'costo',
    'stock',
    'nuevo',
    'promocion',
    'activo',
]


REQUIRED_COLUMNS = {
    'sku',
    'nombre',
    'categoria',
    'precio_detal',
    'precio_mayorista',
    'costo',
    'stock',
}


HEADER_ALIASES = {
    'sku': 'sku',
    'codigo': 'sku',
    'codigo_sku': 'sku',
    'codigo_producto': 'sku',

    'nombre': 'nombre',
    'producto': 'nombre',
    'nombre_producto': 'nombre',

    'categoria': 'categoria',
    'marca': 'marca',

    'descripcion_corta': 'descripcion_corta',
    'descripcion': 'descripcion',

    'precio_detal': 'precio_detal',
    'precio_venta': 'precio_detal',

    'precio_mayorista': 'precio_mayorista',
    'precio_mayor': 'precio_mayorista',

    'costo': 'costo',
    'costo_interno': 'costo',

    'stock': 'stock',
    'existencias': 'stock',
    'cantidad': 'stock',

    'nuevo': 'nuevo',
    'es_nuevo': 'nuevo',

    'promocion': 'promocion',
    'en_promocion': 'promocion',

    'activo': 'activo',
    'es_activo': 'activo',
}


TRUE_VALUES = {
    'SI',
    'SÍ',
    'TRUE',
    'VERDADERO',
    '1',
    'YES',
}


FALSE_VALUES = {
    'NO',
    'FALSE',
    'FALSO',
    '0',
}


# ============================================================
# EXCEPCIÓN
# ============================================================

class ProductImportError(Exception):
    pass


# ============================================================
# FORMULARIO
# ============================================================

class ProductImportUploadForm(forms.Form):

    file = forms.FileField(
        label='Archivo de productos',
        help_text='Formatos permitidos: .xlsx y .csv.',
        widget=forms.ClearableFileInput(
            attrs={
                'accept': '.xlsx,.csv',
            }
        ),
    )

    def clean_file(self):

        uploaded_file = self.cleaned_data['file']

        extension = (
            Path(uploaded_file.name)
            .suffix
            .lower()
        )

        if extension not in ALLOWED_EXTENSIONS:

            raise forms.ValidationError(
                'El archivo debe ser Excel (.xlsx) o CSV (.csv).'
            )

        if uploaded_file.size > MAX_FILE_SIZE:

            raise forms.ValidationError(
                'El archivo supera el tamaño máximo permitido de 15 MB.'
            )

        return uploaded_file


# ============================================================
# NORMALIZACIÓN
# ============================================================

def _normalize_header(value):

    if value is None:
        return ''

    text = str(value).strip()

    text = unicodedata.normalize(
        'NFKD',
        text,
    )

    text = ''.join(
        character
        for character in text
        if not unicodedata.combining(character)
    )

    text = (
        text
        .lower()
        .strip()
    )

    text = re.sub(
        r'[^a-z0-9]+',
        '_',
        text,
    )

    text = text.strip('_')

    return HEADER_ALIASES.get(
        text,
        text,
    )


def _clean_text(value):

    if value is None:
        return ''

    return str(value).strip()


# ============================================================
# DECIMALES
# ============================================================

def _parse_decimal(
    value,
    field_label,
):

    if (
        value is None
        or str(value).strip() == ''
    ):

        return (
            None,
            f'{field_label} es obligatorio.',
        )

    if isinstance(value, bool):

        return (
            None,
            f'{field_label} debe ser un número.',
        )

    if isinstance(
        value,
        (
            int,
            float,
            Decimal,
        )
    ):

        try:

            decimal_value = Decimal(
                str(value)
            )

        except (
            InvalidOperation,
            ValueError,
        ):

            return (
                None,
                f'{field_label} es inválido.',
            )

    else:

        raw_value = str(value).strip()

        if not re.fullmatch(
            r'\d+(?:\.\d{1,2})?',
            raw_value,
        ):

            return (
                None,
                (
                    f'{field_label} debe ser numérico. '
                    f'Use, por ejemplo, 35000.'
                ),
            )

        try:

            decimal_value = Decimal(
                raw_value
            )

        except InvalidOperation:

            return (
                None,
                f'{field_label} es inválido.',
            )

    if decimal_value < 0:

        return (
            None,
            f'{field_label} no puede ser negativo.',
        )

    if decimal_value >= Decimal('10000000000'):

        return (
            None,
            (
                f'{field_label} supera '
                f'el valor máximo permitido.'
            ),
        )

    return (
        decimal_value,
        None,
    )


# ============================================================
# STOCK
# ============================================================

def _parse_stock(value):

    if (
        value is None
        or str(value).strip() == ''
    ):

        return (
            None,
            'Stock es obligatorio.',
        )

    if isinstance(value, bool):

        return (
            None,
            'Stock debe ser un número entero.',
        )

    if isinstance(value, int):

        stock = value

    elif isinstance(value, float):

        if not value.is_integer():

            return (
                None,
                'Stock debe ser un número entero.',
            )

        stock = int(value)

    else:

        raw_value = str(value).strip()

        if not re.fullmatch(
            r'\d+',
            raw_value,
        ):

            return (
                None,
                (
                    'Stock debe ser un número entero '
                    'igual o mayor que cero.'
                ),
            )

        stock = int(raw_value)

    if stock < 0:

        return (
            None,
            'Stock no puede ser negativo.',
        )

    return (
        stock,
        None,
    )


# ============================================================
# BOOLEANOS
# ============================================================

def _parse_boolean(
    value,
    field_label,
):

    if (
        value is None
        or str(value).strip() == ''
    ):

        return (
            None,
            None,
        )

    if isinstance(value, bool):

        return (
            value,
            None,
        )

    normalized = (
        str(value)
        .strip()
        .upper()
    )

    if normalized in TRUE_VALUES:

        return (
            True,
            None,
        )

    if normalized in FALSE_VALUES:

        return (
            False,
            None,
        )

    return (
        None,
        f'{field_label} debe usar SI o NO.',
    )


# ============================================================
# LEER EXCEL
# ============================================================

def _read_xlsx(uploaded_file):

    uploaded_file.seek(0)

    try:

        workbook = load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True,
        )

    except Exception as exc:

        raise ProductImportError(
            (
                'No fue posible leer el archivo Excel. '
                'Verifica que sea un archivo .xlsx válido.'
            )
        ) from exc

    active_sheet = workbook.active

    if active_sheet is None:

        workbook.close()

        raise ProductImportError(
            'El archivo Excel no contiene una hoja activa.'
        )

    worksheet = cast(
        Any,
        active_sheet,
    )

    rows_iterator = worksheet.iter_rows(
        values_only=True,
    )

    try:

        raw_headers = next(
            rows_iterator
        )

    except StopIteration as exc:

        workbook.close()

        raise ProductImportError(
            'El archivo Excel está vacío.'
        ) from exc

    headers = [
        _normalize_header(value)
        for value in raw_headers
    ]

    rows = []

    for row_number, values in enumerate(
        rows_iterator,
        start=2,
    ):

        if all(
            (
                value is None
                or str(value).strip() == ''
            )
            for value in values
        ):

            continue

        row_data = {}

        for index, header in enumerate(
            headers
        ):

            if not header:
                continue

            if index < len(values):

                row_data[header] = (
                    values[index]
                )

            else:

                row_data[header] = None

        rows.append({
            'row_number': row_number,
            'data': row_data,
        })

        if len(rows) > MAX_ROWS:

            workbook.close()

            raise ProductImportError(
                (
                    'El archivo supera el máximo '
                    f'de {MAX_ROWS:,} filas.'
                )
            )

    workbook.close()

    return (
        headers,
        rows,
    )


# ============================================================
# LEER CSV
# ============================================================

def _decode_csv(raw_bytes):

    encodings = [
        'utf-8-sig',
        'utf-8',
        'cp1252',
        'latin-1',
    ]

    for encoding in encodings:

        try:

            return raw_bytes.decode(
                encoding
            )

        except UnicodeDecodeError:

            continue

    raise ProductImportError(
        (
            'No fue posible leer la codificación '
            'del archivo CSV.'
        )
    )


def _read_csv(uploaded_file):

    uploaded_file.seek(0)

    raw_bytes = uploaded_file.read()

    text = _decode_csv(
        raw_bytes
    )

    sample = text[:4096]

    try:

        dialect = (
            csv.Sniffer()
            .sniff(
                sample,
                delimiters=',;\t',
            )
        )

        delimiter = dialect.delimiter

    except csv.Error:

        delimiter = ','

    stream = io.StringIO(
        text
    )

    reader = csv.reader(
        stream,
        delimiter=delimiter,
    )

    try:

        raw_headers = next(
            reader
        )

    except StopIteration as exc:

        raise ProductImportError(
            'El archivo CSV está vacío.'
        ) from exc

    headers = [
        _normalize_header(value)
        for value in raw_headers
    ]

    rows = []

    for row_number, values in enumerate(
        reader,
        start=2,
    ):

        if all(
            not str(value).strip()
            for value in values
        ):

            continue

        row_data = {}

        for index, header in enumerate(
            headers
        ):

            if not header:
                continue

            if index < len(values):

                row_data[header] = (
                    values[index]
                )

            else:

                row_data[header] = ''

        rows.append({
            'row_number': row_number,
            'data': row_data,
        })

        if len(rows) > MAX_ROWS:

            raise ProductImportError(
                (
                    'El archivo supera el máximo '
                    f'de {MAX_ROWS:,} filas.'
                )
            )

    return (
        headers,
        rows,
    )


# ============================================================
# LEER ARCHIVO SUBIDO
# ============================================================

def _read_uploaded_file(
    uploaded_file
):

    extension = (
        Path(uploaded_file.name)
        .suffix
        .lower()
    )

    if extension == '.xlsx':

        return _read_xlsx(
            uploaded_file
        )

    if extension == '.csv':

        return _read_csv(
            uploaded_file
        )

    raise ProductImportError(
        'Formato de archivo no compatible.'
    )


# ============================================================
# VALIDAR ENCABEZADOS
# ============================================================

def _validate_headers(
    headers
):

    problems = []

    normalized_headers = [
        header
        for header in headers
        if header
    ]

    duplicates = [
        header
        for header, count
        in Counter(
            normalized_headers
        ).items()
        if count > 1
    ]

    if duplicates:

        problems.append(
            (
                'Hay columnas repetidas: '
                + ', '.join(duplicates)
                + '.'
            )
        )

    missing = sorted(
        REQUIRED_COLUMNS
        - set(normalized_headers)
    )

    if missing:

        problems.append(
            (
                'Faltan columnas obligatorias: '
                + ', '.join(missing)
                + '.'
            )
        )

    return problems


# ============================================================
# VALIDACIÓN PRINCIPAL
# ============================================================

def validate_product_file(
    uploaded_file
):

    headers, raw_rows = (
        _read_uploaded_file(
            uploaded_file
        )
    )

    header_errors = (
        _validate_headers(
            headers
        )
    )

    if header_errors:

        return {
            'filename': uploaded_file.name,
            'file_errors': header_errors,
            'total_rows': 0,
            'valid_count': 0,
            'error_count': 0,
            'warning_count': 0,
            'create_count': 0,
            'update_count': 0,
            'issue_rows': [],
            'valid_preview': [],
            'importable_rows': [],
        }

    category_map = {
        category.name.casefold(): category
        for category
        in Category.objects.all()
    }

    brand_map = {
        brand.name.casefold(): brand
        for brand
        in Brand.objects.all()
    }

    product_map = {
        product.sku.casefold(): product
        for product
        in Product.objects.all()
    }

    sku_counts = Counter()

    for row in raw_rows:

        sku = _clean_text(
            row[
                'data'
            ].get(
                'sku'
            )
        )

        if sku:

            sku_counts[
                sku.casefold()
            ] += 1

    validated_rows = []

    for row in raw_rows:

        row_number = (
            row[
                'row_number'
            ]
        )

        data = (
            row[
                'data'
            ]
        )

        errors = []
        warnings = []

        # ====================================================
        # SKU
        # ====================================================

        sku = _clean_text(
            data.get(
                'sku'
            )
        )

        if not sku:

            errors.append(
                'SKU es obligatorio.'
            )

        elif len(sku) > 60:

            errors.append(
                'SKU supera 60 caracteres.'
            )

        elif (
            sku_counts[
                sku.casefold()
            ]
            > 1
        ):

            errors.append(
                (
                    'SKU repetido dentro '
                    'del archivo.'
                )
            )

        # ====================================================
        # NOMBRE
        # ====================================================

        name = _clean_text(
            data.get(
                'nombre'
            )
        )

        if not name:

            errors.append(
                'Nombre es obligatorio.'
            )

        elif len(name) > 180:

            errors.append(
                'Nombre supera 180 caracteres.'
            )

        # ====================================================
        # CATEGORÍA
        # ====================================================

        category_name = _clean_text(
            data.get(
                'categoria'
            )
        )

        if not category_name:

            errors.append(
                'Categoría es obligatoria.'
            )

        elif (
            category_name.casefold()
            not in category_map
        ):

            errors.append(
                (
                    f'La categoría "{category_name}" '
                    'no existe.'
                )
            )

        # ====================================================
        # MARCA
        # ====================================================

        brand_name = _clean_text(
            data.get(
                'marca'
            )
        )

        if (
            brand_name
            and (
                brand_name.casefold()
                not in brand_map
            )
        ):

            errors.append(
                (
                    f'La marca "{brand_name}" '
                    'no existe.'
                )
            )

        # ====================================================
        # DESCRIPCIÓN CORTA
        # ====================================================

        short_description = _clean_text(
            data.get(
                'descripcion_corta'
            )
        )

        if len(
            short_description
        ) > 280:

            errors.append(
                (
                    'Descripción corta supera '
                    '280 caracteres.'
                )
            )

        # ====================================================
        # PRECIO DETAL
        # ====================================================

        retail_price, error = (
            _parse_decimal(
                data.get(
                    'precio_detal'
                ),
                'Precio detal',
            )
        )

        if error:

            errors.append(
                error
            )

        # ====================================================
        # PRECIO MAYORISTA
        # ====================================================

        wholesale_price, error = (
            _parse_decimal(
                data.get(
                    'precio_mayorista'
                ),
                'Precio mayorista',
            )
        )

        if error:

            errors.append(
                error
            )

        # ====================================================
        # COSTO
        # ====================================================

        cost, error = (
            _parse_decimal(
                data.get(
                    'costo'
                ),
                'Costo',
            )
        )

        if error:

            errors.append(
                error
            )

        # ====================================================
        # STOCK
        # ====================================================

        stock, error = (
            _parse_stock(
                data.get(
                    'stock'
                )
            )
        )

        if error:

            errors.append(
                error
            )

        # ====================================================
        # NUEVO
        # ====================================================

        is_new, error = (
            _parse_boolean(
                data.get(
                    'nuevo'
                ),
                'Nuevo',
            )
        )

        if error:

            errors.append(
                error
            )

        # ====================================================
        # PROMOCIÓN
        # ====================================================

        is_on_sale, error = (
            _parse_boolean(
                data.get(
                    'promocion'
                ),
                'Promoción',
            )
        )

        if error:

            errors.append(
                error
            )

        # ====================================================
        # ACTIVO
        # ====================================================

        is_active, error = (
            _parse_boolean(
                data.get(
                    'activo'
                ),
                'Activo',
            )
        )

        if error:

            errors.append(
                error
            )

        # ====================================================
        # ACCIÓN PREVISTA
        # ====================================================

        existing_product = None

        if sku:

            existing_product = (
                product_map.get(
                    sku.casefold()
                )
            )

        if existing_product:

            action = 'ACTUALIZAR'

        else:

            action = 'CREAR'

        # ====================================================
        # ADVERTENCIAS
        # ====================================================

        if (
            not short_description
            and not errors
        ):

            warnings.append(
                (
                    'Producto sin '
                    'descripción corta.'
                )
            )

        if (
            not brand_name
            and not errors
        ):

            warnings.append(
                'Producto sin marca.'
            )

        # ====================================================
        # ESTADO
        # ====================================================

        if errors:

            status = 'ERROR'

        elif warnings:

            status = 'ADVERTENCIA'

        else:

            status = 'VÁLIDO'

        validated_rows.append({
            'row_number': row_number,
            'sku': sku,
            'name': name,
            'category': category_name,
            'brand': brand_name,
            'action': action,
            'status': status,
            'errors': errors,
            'warnings': warnings,
            'cleaned': {
                'sku': sku,
                'name': name,
                'category_name': (
                    category_name
                ),
                'brand_name': (
                    brand_name
                ),
                'short_description': (
                    short_description
                ),
                'description': (
                    _clean_text(
                        data.get(
                            'descripcion'
                        )
                    )
                ),
                'retail_price': (
                    retail_price
                ),
                'wholesale_price': (
                    wholesale_price
                ),
                'cost': cost,
                'stock': stock,
                'is_new': (
                    is_new
                ),
                'is_on_sale': (
                    is_on_sale
                ),
                'is_active': (
                    is_active
                ),
            },
        })

    importable_rows = [
        row
        for row in validated_rows
        if not row[
            'errors'
        ]
    ]

    error_rows = [
        row
        for row in validated_rows
        if row[
            'errors'
        ]
    ]

    warning_rows = [
        row
        for row in validated_rows
        if (
            not row[
                'errors'
            ]
            and row[
                'warnings'
            ]
        )
    ]

    issue_rows = [
        row
        for row in validated_rows
        if (
            row[
                'errors'
            ]
            or row[
                'warnings'
            ]
        )
    ]

    create_count = sum(
        1
        for row in importable_rows
        if row[
            'action'
        ] == 'CREAR'
    )

    update_count = sum(
        1
        for row in importable_rows
        if row[
            'action'
        ] == 'ACTUALIZAR'
    )

    return {
        'filename': (
            uploaded_file.name
        ),
        'file_errors': [],
        'total_rows': (
            len(
                validated_rows
            )
        ),
        'valid_count': (
            len(
                importable_rows
            )
        ),
        'error_count': (
            len(
                error_rows
            )
        ),
        'warning_count': (
            len(
                warning_rows
            )
        ),
        'create_count': (
            create_count
        ),
        'update_count': (
            update_count
        ),
        'issue_rows': (
            issue_rows
        ),
        'valid_preview': (
            importable_rows[:50]
        ),
        'importable_rows': (
            importable_rows
        ),
    }


# ============================================================
# SLUG ÚNICO
# ============================================================

def _build_unique_product_slug(
    name,
    sku,
):

    base_slug = slugify(
        f'{name}-{sku}'
    )

    if not base_slug:

        base_slug = slugify(
            sku
        )

    if not base_slug:

        base_slug = (
            f'producto-'
            f'{uuid4().hex[:8]}'
        )

    base_slug = (
        base_slug[:190]
    )

    candidate = base_slug

    counter = 2

    while Product.objects.filter(
        slug=candidate
    ).exists():

        suffix = (
            f'-{counter}'
        )

        candidate = (
            base_slug[
                :200 - len(suffix)
            ]
            + suffix
        )

        counter += 1

    return candidate


# ============================================================
# IMPORTAR PRODUCTOS VALIDADOS
# ============================================================

def import_validated_products(
    uploaded_file,
    user,
):

    validation = (
        validate_product_file(
            uploaded_file
        )
    )

    if validation[
        'file_errors'
    ]:

        raise ProductImportError(
            (
                'El archivo contiene errores '
                'de estructura y no puede '
                'ser importado.'
            )
        )

    rows = (
        validation[
            'importable_rows'
        ]
    )

    if not rows:

        raise ProductImportError(
            (
                'El archivo no contiene '
                'filas válidas para importar.'
            )
        )

    import_reference = (
        'IMPORT-'
        + timezone.now().strftime(
            '%Y%m%d-%H%M%S'
        )
        + '-'
        + uuid4().hex[:6].upper()
    )

    created_count = 0
    updated_count = 0
    movement_count = 0

    with transaction.atomic():

        category_map = {
            category.name.casefold(): category
            for category
            in Category.objects.all()
        }

        brand_map = {
            brand.name.casefold(): brand
            for brand
            in Brand.objects.all()
        }

        for row in rows:

            data = (
                row[
                    'cleaned'
                ]
            )

            sku = (
                data[
                    'sku'
                ]
            )

            category = (
                category_map[
                    data[
                        'category_name'
                    ].casefold()
                ]
            )

            brand = None

            if data[
                'brand_name'
            ]:

                brand = (
                    brand_map[
                        data[
                            'brand_name'
                        ].casefold()
                    ]
                )

            product = (
                Product.objects
                .select_for_update()
                .filter(
                    sku__iexact=sku
                )
                .first()
            )

            # =================================================
            # CREAR PRODUCTO
            # =================================================

            if product is None:

                product = Product(
                    name=(
                        data[
                            'name'
                        ]
                    ),
                    slug=(
                        _build_unique_product_slug(
                            data[
                                'name'
                            ],
                            sku,
                        )
                    ),
                    sku=sku,
                    category=category,
                    brand=brand,
                    short_description=(
                        data[
                            'short_description'
                        ]
                    ),
                    description=(
                        data[
                            'description'
                        ]
                    ),
                    retail_price=(
                        data[
                            'retail_price'
                        ]
                    ),
                    wholesale_price=(
                        data[
                            'wholesale_price'
                        ]
                    ),
                    cost=(
                        data[
                            'cost'
                        ]
                    ),
                    stock=0,
                    is_new=(
                        data[
                            'is_new'
                        ]
                        if (
                            data[
                                'is_new'
                            ]
                            is not None
                        )
                        else False
                    ),
                    is_on_sale=(
                        data[
                            'is_on_sale'
                        ]
                        if (
                            data[
                                'is_on_sale'
                            ]
                            is not None
                        )
                        else False
                    ),
                    is_active=(
                        data[
                            'is_active'
                        ]
                        if (
                            data[
                                'is_active'
                            ]
                            is not None
                        )
                        else True
                    ),
                )

                product.full_clean()

                product.save()

                created_count += 1

                target_stock = (
                    data[
                        'stock'
                    ]
                )

                if target_stock > 0:

                    InventoryMovement.objects.create(
                        product=product,
                        movement_type=(
                            InventoryMovement
                            .MovementType
                            .ENTRY
                        ),
                        quantity=(
                            target_stock
                        ),
                        reason=(
                            'Stock inicial desde '
                            'importación masiva'
                        ),
                        reference=(
                            import_reference
                        ),
                        created_by=(
                            user
                        ),
                    )

                    movement_count += 1

                continue

            # =================================================
            # ACTUALIZAR PRODUCTO EXISTENTE
            # =================================================

            previous_stock = (
                product.stock
            )

            target_stock = (
                data[
                    'stock'
                ]
            )

            product.name = (
                data[
                    'name'
                ]
            )

            product.category = (
                category
            )

            product.brand = (
                brand
            )

            product.short_description = (
                data[
                    'short_description'
                ]
            )

            product.description = (
                data[
                    'description'
                ]
            )

            product.retail_price = (
                data[
                    'retail_price'
                ]
            )

            product.wholesale_price = (
                data[
                    'wholesale_price'
                ]
            )

            product.cost = (
                data[
                    'cost'
                ]
            )

            if (
                data[
                    'is_new'
                ]
                is not None
            ):

                product.is_new = (
                    data[
                        'is_new'
                    ]
                )

            if (
                data[
                    'is_on_sale'
                ]
                is not None
            ):

                product.is_on_sale = (
                    data[
                        'is_on_sale'
                    ]
                )

            if (
                data[
                    'is_active'
                ]
                is not None
            ):

                product.is_active = (
                    data[
                        'is_active'
                    ]
                )

            product.full_clean()

            product.save(
                update_fields=[
                    'name',
                    'category',
                    'brand',
                    'short_description',
                    'description',
                    'retail_price',
                    'wholesale_price',
                    'cost',
                    'is_new',
                    'is_on_sale',
                    'is_active',
                    'updated_at',
                ]
            )

            updated_count += 1

            # =================================================
            # AJUSTAR STOCK
            # =================================================

            difference = (
                target_stock
                - previous_stock
            )

            if difference > 0:

                InventoryMovement.objects.create(
                    product=product,
                    movement_type=(
                        InventoryMovement
                        .MovementType
                        .ADJUSTMENT_IN
                    ),
                    quantity=(
                        difference
                    ),
                    reason=(
                        'Ajuste de stock desde '
                        'importación masiva'
                    ),
                    reference=(
                        import_reference
                    ),
                    created_by=(
                        user
                    ),
                )

                movement_count += 1

            elif difference < 0:

                InventoryMovement.objects.create(
                    product=product,
                    movement_type=(
                        InventoryMovement
                        .MovementType
                        .ADJUSTMENT_OUT
                    ),
                    quantity=(
                        abs(
                            difference
                        )
                    ),
                    reason=(
                        'Ajuste de stock desde '
                        'importación masiva'
                    ),
                    reference=(
                        import_reference
                    ),
                    created_by=(
                        user
                    ),
                )

                movement_count += 1

    return {
        'reference': (
            import_reference
        ),
        'created_count': (
            created_count
        ),
        'updated_count': (
            updated_count
        ),
        'movement_count': (
            movement_count
        ),
        'skipped_error_count': (
            validation[
                'error_count'
            ]
        ),
    }


# ============================================================
# CREAR PLANTILLA EXCEL
# ============================================================

def build_product_import_template():

    workbook = Workbook()

    active_sheet = (
        workbook.active
    )

    if active_sheet is None:

        worksheet = (
            workbook.create_sheet(
                title='Productos'
            )
        )

    else:

        worksheet = cast(
            Worksheet,
            active_sheet,
        )

        worksheet.title = (
            'Productos'
        )

    headers = [
        'sku',
        'nombre',
        'categoria',
        'marca',
        'descripcion_corta',
        'descripcion',
        'precio_detal',
        'precio_mayorista',
        'costo',
        'stock',
        'nuevo',
        'promocion',
        'activo',
    ]

    worksheet.append(
        headers
    )

    worksheet.append([
        'PRONTY-001',
        'Producto de ejemplo',
        'Accesorios',
        'SHEIN',
        'Descripción corta del producto',
        'Descripción completa del producto.',
        35000,
        28000,
        13500,
        10,
        'SI',
        'NO',
        'SI',
    ])

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

    column_widths = {
        'A': 20,
        'B': 30,
        'C': 24,
        'D': 20,
        'E': 38,
        'F': 50,
        'G': 18,
        'H': 20,
        'I': 16,
        'J': 14,
        'K': 12,
        'L': 14,
        'M': 12,
    }

    for column, width in (
        column_widths.items()
    ):

        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.freeze_panes = (
        'A2'
    )

    # ========================================================
    # HOJA DE INSTRUCCIONES
    # ========================================================

    instructions = (
        workbook.create_sheet(
            title='Instrucciones'
        )
    )

    instruction_rows = [
        [
            (
                'PLANTILLA DE IMPORTACIÓN '
                'DE PRODUCTOS PRONTY'
            )
        ],
        [''],
        [
            'Regla',
            'Descripción',
        ],
        [
            'SKU',
            (
                'Obligatorio y único. '
                'Máximo 60 caracteres.'
            ),
        ],
        [
            'Nombre',
            (
                'Obligatorio. '
                'Máximo 180 caracteres.'
            ),
        ],
        [
            'Categoría',
            (
                'Debe coincidir con una '
                'categoría existente en Pronty.'
            ),
        ],
        [
            'Marca',
            (
                'Opcional. Si se escribe, '
                'debe existir en Pronty.'
            ),
        ],
        [
            'Descripción corta',
            (
                'Opcional. Máximo '
                '280 caracteres.'
            ),
        ],
        [
            'Descripción',
            'Opcional.',
        ],
        [
            'Precio detal',
            (
                'Número sin símbolo $. '
                'Ejemplo: 35000.'
            ),
        ],
        [
            'Precio mayorista',
            (
                'Número sin símbolo $. '
                'Ejemplo: 28000.'
            ),
        ],
        [
            'Costo',
            (
                'Número sin símbolo $. '
                'Ejemplo: 13500.'
            ),
        ],
        [
            'Stock',
            (
                'Stock DISPONIBLE para vender '
                'en Pronty. No incluya unidades '
                'que ya estén reservadas.'
            ),
        ],
        [
            'Nuevo',
            'Use SI o NO.',
        ],
        [
            'Promoción',
            'Use SI o NO.',
        ],
        [
            'Activo',
            'Use SI o NO.',
        ],
        [''],
        [
            'PRODUCTOS EXISTENTES',
            (
                'Si el SKU ya existe, '
                'se actualizará el producto.'
            ),
        ],
        [
            'Stock existente',
            (
                'El valor del archivo será el '
                'nuevo stock disponible objetivo. '
                'La diferencia quedará registrada '
                'como movimiento de inventario.'
            ),
        ],
        [''],
        [
            'IMÁGENES',
            (
                'Las imágenes NO se insertan '
                'dentro del Excel.'
            ),
        ],
        [
            'Una imagen',
            (
                'Nombre el archivo igual que '
                'el SKU. Ejemplo: P82.jpg'
            ),
        ],
        [
            'Varias imágenes',
            (
                'Use P82-1.jpg, P82-2.jpg, '
                'P82-3.jpg, etc.'
            ),
        ],
        [
            'Imagen principal',
            (
                'Cuando existen varias, '
                'la terminada en -1 será '
                'la imagen principal.'
            ),
        ],
        [
            'ZIP de imágenes',
            (
                'Las fotografías se cargarán '
                'en un ZIP separado.'
            ),
        ],
    ]

    for row in instruction_rows:

        instructions.append(
            row
        )

    instructions[
        'A1'
    ].font = Font(
        bold=True,
        size=14,
    )

    instructions[
        'A3'
    ].font = Font(
        bold=True
    )

    instructions[
        'B3'
    ].font = Font(
        bold=True
    )

    instructions.column_dimensions[
        'A'
    ].width = 30

    instructions.column_dimensions[
        'B'
    ].width = 85

    return workbook